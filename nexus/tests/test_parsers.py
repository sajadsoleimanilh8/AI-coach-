from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from nexus.core.exceptions import UnsupportedDocumentType
from nexus.rag.parsers import parse_document


@pytest.mark.asyncio
async def test_parse_txt() -> None:
    text = await parse_document(b"hello world", "txt")
    assert text == "hello world"


@pytest.mark.asyncio
async def test_parse_md() -> None:
    text = await parse_document(b"# Heading\n\nBody text.", "md")
    assert "# Heading" in text


@pytest.mark.asyncio
async def test_parse_code() -> None:
    text = await parse_document(b"def foo():\n    return 1\n", "code")
    assert "def foo" in text


@pytest.mark.asyncio
async def test_parse_csv() -> None:
    raw = b"name,age\nAlice,30\nBob,25\n"
    text = await parse_document(raw, "csv")
    lines = text.splitlines()
    assert lines[0] == "name: Alice, age: 30"
    assert lines[1] == "name: Bob, age: 25"


@pytest.mark.asyncio
async def test_parse_csv_empty_is_empty_string() -> None:
    assert await parse_document(b"", "csv") == ""


@pytest.mark.asyncio
async def test_parse_docx() -> None:
    doc = Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    buffer = io.BytesIO()
    doc.save(buffer)

    text = await parse_document(buffer.getvalue(), "docx")

    assert "First paragraph." in text
    assert "Second paragraph." in text


@pytest.mark.asyncio
async def test_parse_pdf() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)

    text = await parse_document(buffer.getvalue(), "pdf")
    assert text == ""


@pytest.mark.asyncio
async def test_unsupported_source_type_raises() -> None:
    with pytest.raises(UnsupportedDocumentType):
        await parse_document(b"data", "exe")


@pytest.mark.asyncio
async def test_source_type_is_case_insensitive() -> None:
    text = await parse_document(b"hello", "TXT")
    assert text == "hello"
